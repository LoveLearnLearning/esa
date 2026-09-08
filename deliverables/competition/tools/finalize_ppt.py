#!/usr/bin/env python3
"""Apply evidence-safe text fixes to the competition PPTX and verify the result."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "04—作品方案/星知智链_作品方案_20260908_修订版.pptx"
MAX_BYTES = 100 * 1024 * 1024


def replace_runs(shape, replacements: dict[str, str]) -> None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            for old, new in replacements.items():
                run.text = run.text.replace(old, new)


def replace_full_text(shape, text: str) -> None:
    text_frame = shape.text_frame
    first_paragraph = text_frame.paragraphs[0]
    first_run = first_paragraph.runs[0] if first_paragraph.runs else first_paragraph.add_run()
    first_run.text = text
    for run in first_paragraph.runs[1:]:
        run.text = ""
    for paragraph in list(text_frame.paragraphs[1:]):
        text_frame._txBody.remove(paragraph._p)


def replace_paragraph_text(shape, paragraph_index: int, text: str) -> None:
    paragraph = shape.text_frame.paragraphs[paragraph_index]
    first_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first_run.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_paragraphs(shape, texts: tuple[str, ...]) -> None:
    text_frame = shape.text_frame
    while len(text_frame.paragraphs) < len(texts):
        text_frame.add_paragraph()
    for index, text in enumerate(texts):
        replace_paragraph_text(shape, index, text)
    for paragraph in list(text_frame.paragraphs[len(texts) :]):
        text_frame._txBody.remove(paragraph._p)


def apply_fixes(presentation: Presentation) -> None:
    slides = presentation.slides

    replace_full_text(
        slides[5].shapes[10],
        "通用问答能给答案，却很难留下下一次学习、教学或科研行动所需的证据。",
    )
    replace_paragraphs(
        slides[9].shapes[7],
        ("基模已完成 LoRA 适配，自建 RAG 系统", "覆盖学习、教学与科研场景"),
    )

    replace_runs(slides[12].shapes[4], {"推理期闸门将": "推理期闸门实验将"})
    replace_full_text(
        slides[12].shapes[42],
        "需要 output_hidden_states；方向与 Adapter 绑定；尚未接入默认线上路径。",
    )

    replace_paragraphs(
        slides[33].shapes[19],
        ("Monaco", "代码块编辑、复制和运行入口。"),
    )
    replace_paragraphs(
        slides[33].shapes[23],
        (
            "Sandbox",
            "Bubblewrap 设计支持 CPU、内存、网络边界；默认关闭，启用需核对主机条件。",
        ),
    )
    replace_paragraphs(
        slides[33].shapes[24],
        ("Audit", "执行和关键教学操作保持可追踪。"),
    )
    replace_paragraphs(
        slides[33].shapes[25],
        ("LSP", "clangd / pyright 可接入；服务端不可用时降级为本地补全。"),
    )

    replace_full_text(slides[34].shapes[7], "AI 不伪造数据、来源或研究结论")
    replace_full_text(
        slides[41].shapes[7],
        "至少 2 名真实目标用户的身份、时间、频次、场景与反馈须由本人确认",
    )
    replace_full_text(slides[43].shapes[20], "通过 · 24.71 s")
    replace_paragraph_text(slides[45].shapes[4], 0, "36 / TEAM")
    replace_full_text(
        slides[45].shapes[7],
        "成员与指导教师信息须与审核通过的报名表逐项一致",
    )

    for slide in slides:
        for shape in walk_shapes(slide.shapes):
            text = getattr(shape, "text", "")
            if "PPT模板" in text or "1ppt.com" in text:
                replace_full_text(shape, "")


def walk_shapes(shapes) -> Iterable:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def all_text(presentation: Presentation) -> str:
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in walk_shapes(slide.shapes)
        if hasattr(shape, "text")
    )


def verify(output: Path) -> None:
    presentation = Presentation(output)
    if len(presentation.slides) != 47:
        raise RuntimeError(f"Expected 47 slides, found {len(presentation.slides)}")
    if output.stat().st_size >= MAX_BYTES:
        raise RuntimeError(f"PPTX exceeds 100 MB: {output.stat().st_size} bytes")

    text = all_text(presentation)
    forbidden = (
        "PPT模板",
        "1ppt.com",
        "证据证据",
        "('Monaco'",
        "('Sandbox'",
        "('Audit'",
        "('LSP'",
        "20余名",
        "共20",
        "_x000B_",
    )
    present = [value for value in forbidden if value in text]
    if present:
        raise RuntimeError(f"Forbidden text remains: {present}")

    required = (
        "星知智链",
        "XH-202620",
        "通过. 20.95 s",
        "通过 · 22.80s",
        "通过 · 24.71 s",
        "尚未接入默认线上路径",
        "至少 2 名真实目标用户",
        "36 / TEAM",
    )
    missing = [value for value in required if value not in text]
    if missing:
        raise RuntimeError(f"Required text missing: {missing}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Source PPTX")
    parser.add_argument("output", type=Path, nargs="?", default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not args.input.exists():
        raise FileNotFoundError(args.input)
    presentation = Presentation(args.input)
    apply_fixes(presentation)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(args.output)
    verify(args.output)
    print(f"built: {args.output}")
    print(f"slides: 47; size: {args.output.stat().st_size / 1024 / 1024:.2f} MiB")


if __name__ == "__main__":
    main()
