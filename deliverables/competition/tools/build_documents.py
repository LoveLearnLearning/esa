#!/usr/bin/env python3
"""Build editable DOCX files from the competition Markdown sources."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCUMENTS = {
    ROOT / "README.md": ROOT / "星知智链_初赛交付说明.docx",
    ROOT / "提交清单.md": ROOT / "星知智链_提交清单.docx",
    ROOT / "01—参赛信息/README.md": ROOT / "01—参赛信息/参赛信息核对说明.docx",
    ROOT / "02—伦理与安全合规性声明/伦理与安全合规性声明.md": (
        ROOT / "02—伦理与安全合规性声明/伦理与安全合规性声明.docx"
    ),
    ROOT / "03—作品Demo/作品Demo说明.md": ROOT / "03—作品Demo/作品Demo说明.docx",
    ROOT / "04—作品方案/PPT核验说明.md": ROOT / "04—作品方案/PPT核验说明.docx",
    ROOT / "05—作品代码/作品代码与复现说明.md": (
        ROOT / "05—作品代码/作品代码与复现说明.docx"
    ),
    ROOT / "05—作品代码/模型与ServiceID交付确认表.md": (
        ROOT / "05—作品代码/模型与ServiceID交付确认表.docx"
    ),
    ROOT / "06—效果验证报告/效果验证报告.md": (
        ROOT / "06—效果验证报告/效果验证报告.docx"
    ),
    ROOT / "06—效果验证报告/真实用户反馈确认表.md": (
        ROOT / "06—效果验证报告/真实用户反馈确认表.docx"
    ),
    ROOT / "06—效果验证报告/典型问题记录表.md": (
        ROOT / "06—效果验证报告/典型问题记录表.docx"
    ),
    ROOT / "06—效果验证报告/视频Demo脚本.md": (
        ROOT / "06—效果验证报告/视频Demo脚本.docx"
    ),
    ROOT / "07—其他材料/其他材料说明.md": ROOT / "07—其他材料/其他材料说明.docx",
    ROOT / "07—其他材料/工程验证摘要_20260908.md": (
        ROOT / "07—其他材料/工程验证摘要_20260908.docx"
    ),
}

BLUE = "174A7E"
ORANGE = "E75B27"
LIGHT_BLUE = "EAF1F7"
LIGHT_GRAY = "F3F4F6"
MID_GRAY = "D9DEE5"
TEXT = RGBColor(31, 41, 55)

STRUCTURAL = re.compile(
    r"^(#{1,4}\s|```|>\s?|[-*]\s|\d+\.\s|\|.*\||<!--\s*pagebreak\s*-->)"
)
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^]]+\]\([^)]+\))")


def set_run_font(run, east_asia: str = "Microsoft YaHei", latin: str = "Arial") -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, value: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(107, 114, 128)


def configure_document(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color, before, after in (
        ("Title", 25, BLUE, 0, 14),
        ("Heading 1", 17, BLUE, 14, 7),
        ("Heading 2", 13.5, BLUE, 11, 5),
        ("Heading 3", 11.5, ORANGE, 8, 4),
        ("Heading 4", 10.5, ORANGE, 7, 3),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    header = section.header.paragraphs[0]
    header.text = f"星知智链（ESA）  |  XH-202620  |  {title}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(107, 114, 128)
    add_page_number(section.footer.paragraphs[0])

    document.core_properties.title = title
    document.core_properties.subject = "星知智链（ESA）初赛提交材料"
    document.core_properties.keywords = "ESA, XH-202620, 竞赛提交材料"


def add_hyperlink(paragraph, label: str, target: str) -> None:
    relationship = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color, underline))
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run)
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, east_asia="Consolas", latin="Consolas")
            run.font.size = Pt(9.25)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        else:
            label, target = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, target)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)


def add_paragraph(document: Document, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    add_inline(paragraph, text.strip())
    return paragraph


def add_quote(document: Document, text: str) -> None:
    paragraph = add_paragraph(document, text)
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.right_indent = Cm(0.25)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), ORANGE)
    borders.append(left)
    properties.append(borders)


def add_code_block(document: Document, code: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_GRAY)
    properties.append(shading)
    run = paragraph.add_run(code.rstrip())
    set_run_font(run, east_asia="Consolas", latin="Consolas")
    run.font.size = Pt(8.75)


def split_table_row(line: str) -> list[str]:
    return [cell.strip().replace("\\|", "|") for cell in line.strip().strip("|").split("|")]


def is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) > 1 and is_separator_row(rows[1]):
        rows = [rows[0], *rows[2:]]
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            value = values[column_index] if column_index < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, BLUE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_break(document: Document) -> None:
    document.add_section(WD_SECTION.NEW_PAGE)


def render_markdown(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), source.stem)
    document = Document()
    configure_document(document, title)

    index = 0
    first_heading = True
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped == "<!-- pagebreak -->":
            add_page_break(document)
            index += 1
            continue
        if stripped.startswith("```"):
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            add_code_block(document, "\n".join(block))
            continue
        if stripped.startswith("|"):
            table_rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(split_table_row(lines[index]))
                index += 1
            add_table(document, table_rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if first_heading and level == 1:
                paragraph = add_paragraph(document, text, "Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_heading = False
            else:
                add_paragraph(document, text, f"Heading {level}")
            index += 1
            continue
        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            add_quote(document, " ".join(quote_lines))
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            index += 1
            item_lines = [bullet.group(1)]
            while index < len(lines):
                candidate = lines[index].strip()
                if not candidate or STRUCTURAL.match(candidate) or candidate in {"---", "***"}:
                    break
                item_lines.append(candidate)
                index += 1
            add_paragraph(document, " ".join(item_lines), "List Bullet")
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            index += 1
            item_lines = [numbered.group(1)]
            while index < len(lines):
                candidate = lines[index].strip()
                if not candidate or STRUCTURAL.match(candidate) or candidate in {"---", "***"}:
                    break
                item_lines.append(candidate)
                index += 1
            add_paragraph(document, " ".join(item_lines), "List Number")
            continue
        if stripped in {"---", "***"}:
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or STRUCTURAL.match(candidate) or candidate in {"---", "***"}:
                break
            paragraph_lines.append(candidate)
            index += 1
        add_paragraph(document, " ".join(paragraph_lines))

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    reopened = Document(output)
    if not reopened.paragraphs and not reopened.tables:
        raise RuntimeError(f"Generated document is empty: {output}")


def main() -> None:
    for source, output in DOCUMENTS.items():
        if not source.exists():
            raise FileNotFoundError(source)
        render_markdown(source, output)
        print(f"built: {output.relative_to(ROOT.parent.parent)}")


if __name__ == "__main__":
    main()
