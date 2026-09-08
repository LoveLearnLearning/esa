#!/usr/bin/env python3
"""Build the competition technical report DOCX from the checked-in Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "星知智链技术报告.md"
OUTPUT = ROOT / "星知智链技术报告.docx"

BLUE = "174A7E"
ORANGE = "EC6428"
LIGHT_BLUE = "EAF1F7"
LIGHT_GRAY = "F3F4F6"
MID_GRAY = "D9DEE5"
LIGHT_GREEN = "EAF5EE"
LIGHT_YELLOW = "FFF4D6"
LIGHT_RED = "FDECEC"
TEXT = RGBColor(31, 41, 55)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def status_fill(value: str) -> str | None:
    normalized = value.replace(" ", "")
    if any(token in normalized for token in ("失败", "不满足", "未完成")):
        return LIGHT_RED
    if any(
        token in normalized
        for token in ("待补", "待验收", "未达到", "未接入", "不代替", "仍需")
    ):
        return LIGHT_YELLOW
    if any(
        token in normalized
        for token in ("已实现并有仓库证据", "满足并有", "全部通过", "机制满足")
    ):
        return LIGHT_GREEN
    return None


def set_run_font(run, east_asia: str = "Microsoft YaHei", latin: str = "Arial") -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(107, 114, 128)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color, before, after in (
        ("Title", 28, BLUE, 0, 16),
        ("Subtitle", 16, "4B5563", 0, 18),
        ("Heading 1", 18, BLUE, 16, 8),
        ("Heading 2", 14, BLUE, 12, 6),
        ("Heading 3", 11.5, ORANGE, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].font.name = "Arial"
    styles["List Bullet"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["List Number"].font.name = "Arial"
    styles["List Number"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    header = section.header.paragraphs[0]
    header.text = "星知智链（ESA）技术报告  |  XH-202620"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run)
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(107, 114, 128)
    add_page_number(section.footer.paragraphs[0])

    core = document.core_properties
    core.title = "星知智链（ESA）技术报告"
    core.subject = "面向一流学科建设的计算机学科垂类大模型与创新应用开发"
    core.keywords = "ESA, 学科垂类大模型, LoRA, Agent, RAG, 教育智能体"


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run)
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, east_asia="Consolas", latin="Consolas")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)


def add_paragraph(document: Document, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    add_inline(paragraph, text.strip())
    return paragraph


def add_quote(document: Document, text: str) -> None:
    paragraph = add_paragraph(document, text, None)
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.right_indent = Cm(0.3)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), ORANGE)
    borders.append(left)
    p_pr.append(borders)


def add_code_block(document: Document, code: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shading)
    run = paragraph.add_run(code.rstrip())
    set_run_font(run, east_asia="Microsoft YaHei", latin="Consolas")
    run.font.size = Pt(8.5)


def parse_table_row(line: str) -> list[str]:
    return [item.strip() for item in line.strip().strip("|").split("|")]


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    repeat_table_header(table.rows[0])
    status_columns = {
        index
        for index, value in enumerate(rows[0])
        if value.strip() in {"状态", "符合性判断", "结果"}
    }
    for r_index, row in enumerate(rows):
        prevent_row_split(table.rows[r_index])
        for c_index in range(cols):
            cell = table.cell(r_index, c_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = row[c_index] if c_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if r_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if r_index == 0:
                set_cell_shading(cell, BLUE)
            elif c_index in status_columns and (fill := status_fill(value)):
                set_cell_shading(cell, fill)
            elif r_index % 2 == 0:
                set_cell_shading(cell, "F7F9FB")
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_image(document: Document, alt: str, raw_path: str) -> None:
    image_path = (ROOT / raw_path).resolve()
    if not image_path.exists():
        add_paragraph(document, f"[图片缺失：{image_path}]", None)
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.35))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption_run = caption.add_run(f"图：{alt}")
    set_run_font(caption_run)
    caption_run.italic = True
    caption_run.font.size = Pt(9)
    caption_run.font.color.rgb = RGBColor(75, 85, 99)


def add_cover(document: Document, front_lines: list[str]) -> None:
    title = front_lines[0].removeprefix("# ").strip()
    subtitle = front_lines[2].removeprefix("## ").strip()

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(45)

    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run)
    run.font.size = Pt(29)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    accent = document.add_paragraph()
    accent.alignment = WD_ALIGN_PARAGRAPH.CENTER
    accent_run = accent.add_run("━━━━━━━━━━━━━━━━━━━━")
    set_run_font(accent_run)
    accent_run.font.color.rgb = RGBColor.from_string(ORANGE)
    accent_run.font.size = Pt(12)

    p = document.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_run_font(run)
    run.font.size = Pt(16)

    document.add_paragraph().paragraph_format.space_after = Pt(50)
    for line in front_lines[4:8]:
        if not line.strip():
            continue
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, line.replace("  ", "").strip())
        for run in p.runs:
            run.font.size = Pt(11)

    quote = next((line[2:] for line in front_lines if line.startswith("> ")), "")
    if quote:
        document.add_paragraph().paragraph_format.space_after = Pt(24)
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.right_indent = Cm(1.5)
        add_inline(p, quote)
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(75, 85, 99)

    document.add_page_break()


def add_toc(document: Document, headings: list[str]) -> None:
    title = document.add_paragraph(style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("目录")
    for heading in headings:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(heading)
        set_run_font(run)
        run.font.size = Pt(10.5)
        if re.match(r"^\d+\.", heading) or heading.startswith("附录"):
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(BLUE)
    document.add_page_break()


def build_body(document: Document, lines: list[str]) -> None:
    i = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = "".join(
                item[:-2] if item.endswith("  ") else item + " "
                for item in paragraph_buffer
            ).strip()
            add_paragraph(document, text)
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped == "<!-- pagebreak -->":
            flush_paragraph()
            document.add_page_break()
            i += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(document, "\n".join(code_lines))
            i += 1
            continue
        if stripped.startswith("| ") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1]):
            flush_paragraph()
            rows = [parse_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(document, rows)
            continue
        image = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image:
            flush_paragraph()
            add_image(document, image.group(1), image.group(2))
            i += 1
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            add_paragraph(document, stripped[4:], "Heading 2")
            i += 1
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            add_paragraph(document, stripped[3:], "Heading 1")
            i += 1
            continue
        if stripped.startswith("# "):
            flush_paragraph()
            add_paragraph(document, stripped[2:], "Title")
            i += 1
            continue
        if stripped.startswith("> "):
            flush_paragraph()
            add_quote(document, stripped[2:])
            i += 1
            continue
        if re.match(r"^- \[[ xX]\] ", stripped):
            flush_paragraph()
            checked = stripped[3].lower() == "x"
            add_paragraph(document, ("☒ " if checked else "☐ ") + stripped[6:])
            i += 1
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            add_paragraph(document, stripped[2:], "List Bullet")
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            flush_paragraph()
            add_paragraph(document, numbered.group(1), "List Number")
            i += 1
            continue
        paragraph_buffer.append(line)
        i += 1
    flush_paragraph()


def main() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    marker = lines.index("<!-- pagebreak -->")
    front = lines[:marker]
    body = lines[marker + 1 :]
    headings = [line[3:].strip() for line in body if line.startswith("## ")]

    document = Document()
    configure_document(document)
    add_cover(document, front)
    add_toc(document, headings)
    build_body(document, body)

    for section in document.sections[1:]:
        section.start_type = WD_SECTION.NEW_PAGE
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
